import json
import os
import time
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import yaml
import alfworld
from alfworld.agents.environment.alfred_tw_env import AlfredTWEnv
from openai import OpenAI


@dataclass
class StepRecord:
    kind: str
    text: str


@dataclass
class WorldModelState:
    goal: str = ""
    observations: List[Tuple[str, str]] = field(default_factory=list)
    knowledge: Dict[str, Any] = field(default_factory=dict)

    def __post_init__(self):
        if "blacklist" not in self.knowledge:
            self.knowledge["blacklist"] = set()
        if "visited_states" not in self.knowledge:
            self.knowledge["visited_states"] = set()
        if "tried_actions" not in self.knowledge:
            self.knowledge["tried_actions"] = set()
        if "action_state_count" not in self.knowledge:
            self.knowledge["action_state_count"] = {}
        if "seen_self_loops" not in self.knowledge:
            self.knowledge["seen_self_loops"] = set()
        if "filter_principles" not in self.knowledge:
            self.knowledge["filter_principles"] = [
                "Rarity first: [rarity=N] marks historical occurrence frequency. Smaller N means the action is less common and may be a fleeting opportunity, so retain it preferentially. Large-N actions are widely available and need not be selected immediately.",
                "Prefer actions directly related to the current thought that can advance the task goal.",
                "Avoid actions that clearly do not change the situation, such as merely looking or confirming the current state.",
            ]


@dataclass
class StuckState:
    is_stuck: bool = False
    last_obs: str = ""
    repeat_count: int = 0
    stuck_duration: int = 0
    stuck_actions: List[str] = field(default_factory=list)


def build_client() -> OpenAI:
    base_url = os.getenv("OPENAI_BASE_URL", "https://api.chatanywhere.tech/v1")
    api_key = os.getenv("OPENAI_API_KEY", "")
    if not api_key:
        raise ValueError("OPENAI_API_KEY is required")
    return OpenAI(base_url=base_url, api_key=api_key)


def llm(prompt: str, client: OpenAI, model: str, max_tokens: int) -> str:
    for attempt in range(2):
        try:
            messages = [{"role": "user", "content": prompt}]
            response = client.chat.completions.create(
                model=model,
                messages=messages,
                stream=False,
                max_tokens=max_tokens,
                timeout=60.0,
            )
            content = response.choices[0].message.content or ""
            return content.strip()
        except Exception as e:
            print(f"LLM call failed (attempt {attempt + 1}/2): {str(e)[:100]}")
            if attempt == 0:
                time.sleep(2)
                continue
            return ""


def process_ob(ob: str) -> str:
    if ob.startswith("You arrive at loc "):
        ob = ob[ob.find(". ") + 2 :]
    return ob


def get_admissible(info: Optional[Dict]) -> List[str]:
    if isinstance(info, dict):
        ac = info.get("admissible_commands")
        if isinstance(ac, list) and len(ac) > 0 and isinstance(ac[0], list):
            return [a for a in ac[0] if isinstance(a, str)]
    return []


def extract_goal(ob: str) -> str:
    marker = "Your task is to:"
    if marker in ob:
        return ob.split(marker, 1)[1].strip().split("\n")[0]
    return ""


def simulate_action(
    state: WorldModelState, action: str, client: OpenAI, model: str, max_tokens: int
) -> str:
    recent_obs = (
        state.observations[-3:] if len(state.observations) > 3 else state.observations
    )
    history_text = "\n".join(
        [
            f"{i + 1}. Executed: {a}\n   Observation: {o}"
            for i, (a, o) in enumerate(recent_obs)
        ]
    )
    if not history_text:
        history_text = "(initial state; no history)"
    if state.observations:
        full_history_text = "\n".join(
            [
                f"{i + 1}. Executed: {a} -> Observation: {o[:80]}"
                for i, (a, o) in enumerate(state.observations)
            ]
        )
    else:
        full_history_text = "(initial state; no history)"
    prompt_step1 = f"You are a world-model simulator. Given the task goal and historical observations, predict the most likely observation after executing an action.\n\nTask goal: {state.goal}\n\nRecent history:\n{history_text}\n\nAction to predict: {action}\n\nPrediction requirements (important):\n1. Make conservative predictions strictly based on historical observations and the literal effect of the action. Do not invent favorable information to advance the task.\n2. If a location was previously visited and the target object was not observed, revisiting it MUST NOT cause the target to appear without evidence.\n3. A movement (go to) or opening (open) action usually only changes the location or visible container contents.\n   Do not assume that a required object happens to be there unless the history explicitly supports it.\n\nPredict the most likely resulting observation. Output only the predicted observation text, without explanation.\n"
    predicted_obs = llm(prompt_step1, client=client, model=model, max_tokens=max_tokens)
    predicted_obs = predicted_obs.strip()
    prompt_step2 = f'You are a task-progress assessment expert. Given the task goal, the complete execution history so far, and the proposed action, assess how much the action advances the task.\n\nTask goal: {state.goal}\n\nComplete history so far (executed actions and real observations, in order):\n{full_history_text}\n\nProposed action: {action}\nPredicted observation for this action: {predicted_obs}\n\nAssessment criteria:\n1. Use the complete history to identify substeps that are ALREADY COMPLETED (e.g., whether an object has been acquired, processed, or placed).\n2. Do NOT treat completed substeps as pending. Repeating them is going in circles and should be rated "No clear progress".\n3. Progress means advancing an UNFINISHED substep or directly completing the task.\n4. Reaching a location that is a prerequisite for the next necessary substep may count as "Some progress, but insufficient".\n\nOutput exactly one of these formats, with a brief reason:\n- Task completed: [reason]\n- Near completion: [reason]\n- Some progress, but insufficient: [reason]\n- No clear progress: [reason]\n'
    task_progress = llm(prompt_step2, client=client, model=model, max_tokens=max_tokens)
    task_progress = task_progress.strip()
    combined = f"Predicted observation: {predicted_obs}\nTask progress: {task_progress}"
    return combined


def quick_filter_candidates(
    state: WorldModelState,
    all_candidates: List[str],
    target_count: int,
    client: OpenAI,
    model: str,
    max_tokens: int,
    current_think: str = "",
) -> List[str]:
    if len(all_candidates) <= target_count:
        return all_candidates
    counts = state.knowledge.get("action_state_count", {})
    candidates_text = "\n".join(
        [
            f"{i + 1}. {a}  [rarity={counts.get(a, 0)}]"
            for i, a in enumerate(all_candidates)
        ]
    )
    recent = (
        state.observations[-3:] if len(state.observations) > 3 else state.observations
    )
    history_text = (
        "\n".join([f"{a} → {o[:60]}..." for a, o in recent])
        if recent
        else "(initial state)"
    )
    think_section = f"\nCurrent thought: {current_think}\n" if current_think else ""
    principles = state.knowledge.get("filter_principles", [])
    principles_text = (
        "\n".join([f"- {p}" for p in principles]) if principles else "(none)"
    )
    prompt = f"You are an action-filtering expert. Use the task goal, current thought, recent history, and filtering principles to select the most promising {target_count} candidates.\n\nTask goal: {state.goal}\n{think_section}\nRecent history:\n{history_text}\n\nFiltering principles (apply jointly):\n{principles_text}\n\nCandidate actions ([rarity=N] means that the action has appeared N times in history; a smaller N means a rarer action):\n{candidates_text}\n\nApply these principles to select the {target_count} actions most likely to advance the task. Output only their indices, separated by commas, e.g.: 3,7,12,15,18,20,22,25,27,28\n"
    output = llm(prompt, client=client, model=model, max_tokens=max_tokens)
    try:
        selected_indices = [int(x.strip()) - 1 for x in output.strip().split(",")]
        selected = [
            all_candidates[i] for i in selected_indices if 0 <= i < len(all_candidates)
        ]
        if len(selected) >= target_count // 2:
            print(f"  Quick filtering succeeded: selected {len(selected)} candidates")
            return selected[:target_count]
    except Exception as e:
        print(
            f"  Quick-filter parsing failed: {e}, falling back to category-based sampling"
        )
    return sample_diverse_candidates(all_candidates, target_count)


def sample_diverse_candidates(actions: List[str], max_count: int) -> List[str]:
    from collections import defaultdict
    import random

    groups = defaultdict(list)
    for act in actions:
        action_type = act.split()[0].lower()
        groups[action_type].append(act)
    sampled = []
    per_group = max(1, max_count // max(1, len(groups)))
    for group_actions in groups.values():
        sampled.extend(random.sample(group_actions, min(per_group, len(group_actions))))
    if len(sampled) < max_count:
        remaining = [a for a in actions if a not in sampled]
        needed = max_count - len(sampled)
        if remaining:
            sampled.extend(random.sample(remaining, min(needed, len(remaining))))
    return sampled[:max_count]


def select_best_action(
    state: WorldModelState,
    candidates: List[Tuple[str, str]],
    client: OpenAI,
    model: str,
    max_tokens: int,
) -> str:
    if not candidates:
        return ""
    candidates_text = ""
    for i, (action, combined_pred) in enumerate(candidates, 1):
        candidates_text += f"{i}. Action: {action}\n{combined_pred}\n\n"
    prompt = f'You are a path-selection expert. Given the task goal, select the action most likely to complete the task from the candidates and their predictions.\n\nTask goal: {state.goal}\n\nEach candidate has two parts:\n  - Predicted observation: what may be observed after execution\n  - Task progress: how much the action advances the goal (completed / near completion / some progress / no progress)\n\nCandidate actions and predictions:\n{candidates_text}\n\nSelection rules, in priority order:\n1. **Prefer actions rated "Task completed" or "Near completion"**: these most directly advance the task.\n2. If none complete the task, choose "Some progress, but insufficient": at least it moves in the right direction.\n3. Avoid "No clear progress": these repeat the current situation or deviate from the goal.\n\nOutput only the action index (1-{len(candidates)}), without explanation.\n'
    output = llm(prompt, client=client, model=model, max_tokens=max_tokens)
    try:
        choice = int(output.strip())
        if 1 <= choice <= len(candidates):
            return candidates[choice - 1][0]
    except:
        pass
    return candidates[0][0]


def re_review_deadlock(
    state: WorldModelState,
    admissible: List[str],
    stuck_intent: str,
    reason: str,
    client: OpenAI,
    model: str,
    max_tokens: int,
) -> Optional[str]:
    if not admissible:
        return None
    if state.observations:
        full_history = "\n".join(
            [
                f"{i + 1}. Executed: {a} -> Observation: {o[:80]}"
                for i, (a, o) in enumerate(state.observations)
            ]
        )
    else:
        full_history = "(no history)"
    actions_text = "\n".join([f"{i + 1}. {a}" for i, a in enumerate(admissible)])
    prompt = f"You are an error-correction and replanning expert. The agent is currently stuck: {reason}\nIt repeatedly insists on a plan without advancing the task. The plan may be wrong or missing a prerequisite.\n\nTask goal: {state.goal}\n\nComplete history so far (executed actions and real observations, in order):\n{full_history}\n\nIntent the agent keeps proposing but cannot ground: {stuck_intent}\n\n[Currently executable actions] (select only from this list; indices start at 1):\n{actions_text}\n\nChoose between these two cases:\nA. If the intent is correct but lacks a prerequisite (e.g., acquiring an object, opening a container,\n   or reaching a location), select the action that BEST SATISFIES that prerequisite.\nB. If the intent is misguided (e.g., the object is not here or the interaction does not exist),\n   select an action CLEARLY DIFFERENT from the recent repeated attempts to explore another possibility.\n\nOutput only the selected action INDEX (1-{len(admissible)}), without explanation.\n"
    output = llm(prompt, client=client, model=model, max_tokens=max_tokens)
    try:
        choice = int(output.strip().split()[0].rstrip(".,)"))
        if 1 <= choice <= len(admissible):
            return admissible[choice - 1]
    except Exception:
        pass
    return None


def state_fingerprint(admissible: Optional[List[str]]) -> Optional[frozenset]:
    if admissible is None:
        return None
    return frozenset(admissible)


def rank_by_rarity(state: WorldModelState, actions: List[str]) -> List[str]:
    counts = state.knowledge.get("action_state_count", {})
    tried = state.knowledge.get("tried_actions", set())
    NOISE_THRESHOLD = 5
    selectable = []
    noise = []
    for a in actions:
        freq = counts.get(a, 0)
        if freq > NOISE_THRESHOLD and a not in tried:
            noise.append(a)
        else:
            selectable.append(a)
    selectable_sorted = sorted(selectable, key=lambda a: counts.get(a, 0))
    return selectable_sorted + noise


def match_intent_to_action(intent: str, admissible: List[str]) -> Optional[str]:
    if not intent or not admissible:
        return None
    intent_l = intent.strip().lower()
    for a in admissible:
        if a.lower() == intent_l:
            return a
    contained = [a for a in admissible if a.lower() in intent_l]
    if not contained:
        return None
    if len(contained) == 1:
        return contained[0]
    contained.sort(key=len, reverse=True)
    if len(contained[0]) > len(contained[1]):
        return contained[0]
    return None


def update_knowledge_from_observation(
    state: WorldModelState,
    action: str,
    observation: str,
    admissible_before: Optional[List[str]] = None,
    admissible_after: Optional[List[str]] = None,
) -> None:
    fp_before = state_fingerprint(admissible_before)
    fp_after = state_fingerprint(admissible_after)
    if fp_before is not None and fp_after is not None:
        is_self_loop = fp_after == fp_before
        is_revisit = fp_after in state.knowledge["visited_states"]
        key = (fp_before, action)
        tried_action_types = set()
        for a in state.knowledge["tried_actions"]:
            parts = a.split()
            if parts:
                tried_action_types.add(parts[0].lower())
        untried_action_types_in_target = set()
        for a in fp_after:
            parts = a.split()
            if parts:
                action_type = parts[0].lower()
                if action_type not in tried_action_types:
                    untried_action_types_in_target.add(action_type)
        has_untried_opportunity = bool(untried_action_types_in_target)
        if is_self_loop:
            if key in state.knowledge["seen_self_loops"]:
                no_progress = True
            else:
                state.knowledge["seen_self_loops"].add(key)
                no_progress = False
            reason = "Repeated self-loop (world unchanged)"
        else:
            no_progress = is_revisit and (not has_untried_opportunity)
            reason = "Revisited state (backtracking)"
        if no_progress and key not in state.knowledge["blacklist"]:
            state.knowledge["blacklist"].add(key)
            print(
                f"Local pruning: non-progress action [{reason}] '{action}' → {observation[:40]}..."
            )
    else:
        prev_obs = state.observations[-2][1] if len(state.observations) >= 2 else None
        if prev_obs is not None and observation.strip() == prev_obs.strip():
            key = (None, action)
            if key not in state.knowledge["blacklist"]:
                state.knowledge["blacklist"].add(key)
                print(
                    f"Local pruning: ineffective action (observation unchanged) '{action}' → {observation[:40]}..."
                )


def detect_action_loop(last_actions: List[str], window: int = 6) -> bool:
    if len(last_actions) < window:
        return False
    recent = last_actions[-window:]
    if len(recent) >= 4:
        half = len(recent) // 2
        first_half = recent[:half]
        second_half = recent[half : half * 2]
        if first_half == second_half:
            return True
    unique_actions = set(recent)
    if len(unique_actions) <= 2 and len(recent) >= 4:
        return True
    return False


def update_stuck_state(stuck_state: StuckState, last_actions: List[str]) -> None:
    if detect_action_loop(last_actions, window=6):
        if not stuck_state.is_stuck:
            stuck_state.is_stuck = True
            stuck_state.stuck_duration = 4
            stuck_state.stuck_actions = list(set(last_actions[-6:]))
            print(
                f"Action cycle detected; forcing diverse exploration for the next {stuck_state.stuck_duration} steps"
            )
            print(f"   Cyclic actions: {stuck_state.stuck_actions[:3]}...")
    if stuck_state.is_stuck:
        stuck_state.stuck_duration -= 1
        if stuck_state.stuck_duration <= 0:
            stuck_state.is_stuck = False
            stuck_state.stuck_actions.clear()
            print("Stuck mode ended")


def filter_actions_if_stuck(
    actions: List[str], stuck_state: StuckState, last_actions: List[str]
) -> List[str]:
    if not stuck_state.is_stuck:
        return actions
    filtered = [a for a in actions if a not in stuck_state.stuck_actions]
    if filtered:
        print(
            f"Stuck filter: excluding cyclic actions; remaining {len(filtered)}/{len(actions)} candidates"
        )
        return filtered
    else:
        print(f"Warning: stuck filtering removed all actions; using the original list")
        return actions


def run_episode(
    env,
    init_prompt: str,
    ob: str,
    info: Dict,
    client: OpenAI,
    model: str,
    max_tokens: int,
    max_steps: int = 50,
) -> Tuple[float, bool, List[StepRecord], int]:
    prompt = ""
    trajectory: List[StepRecord] = []
    goal = extract_goal(ob)
    state = WorldModelState(goal=goal)
    stuck_state = StuckState()
    last_actions: List[str] = []
    last_intent: str = ""
    intent_repeat_count: int = 0
    optimistic_streak: int = 0
    DEADLOCK_K = 3
    llm_call_count = 0
    print(f"Task goal: {goal}")
    print(f"Maximum steps: {max_steps}\n")
    for step in range(1, max_steps + 1):
        print(f"--- Step {step}/{max_steps} ---")
        admissible = get_admissible(info)
        print(f"Number of executable actions: {len(admissible)}")
        if not admissible:
            print("Warning: no executable actions; ending episode")
            break
        asc = state.knowledge["action_state_count"]
        for a in admissible:
            asc[a] = asc.get(a, 0) + 1
        think_prompt = (
            init_prompt
            + prompt
            + "\nFirst output one reasoning line beginning with 'think:', then one line beginning with 'action:' specifying the concrete action you intend to execute immediately."
        )
        think_output = llm(
            think_prompt, client=client, model=model, max_tokens=max_tokens
        )
        llm_call_count += 1
        think_text, intent_action = ("", "")
        for line in think_output.splitlines():
            s = line.strip().lstrip(">-*#• ").strip()
            if s.lower().startswith("think:"):
                think_text = s[6:].strip()
            elif s.lower().startswith("action:"):
                intent_action = s[7:].strip()
        if not think_text:
            think_text = think_output.strip()
        trajectory.append(StepRecord(kind="think", text=think_text))
        print(f"Thought: {think_text}")
        blacklist = state.knowledge["blacklist"]
        cur_fp = state_fingerprint(admissible)
        filtered_admissible = [
            act for act in admissible if (cur_fp, act) not in blacklist
        ]
        if not filtered_admissible:
            print("Warning: all actions are blacklisted; using the original list")
            filtered_admissible = admissible
        elif len(filtered_admissible) < len(admissible):
            print(
                f"Local pruning filter: {len(admissible)} → {len(filtered_admissible)}"
            )
        filtered_admissible = filter_actions_if_stuck(
            filtered_admissible, stuck_state, last_actions
        )
        same_intent_as_last = bool(intent_action) and intent_action == last_intent
        selected_action = None
        if not stuck_state.is_stuck:
            selected_action = match_intent_to_action(intent_action, filtered_admissible)
        took_intent = bool(selected_action)
        is_optimistic = False
        deadlock_reason = ""
        if not took_intent:
            if same_intent_as_last and intent_repeat_count + 1 >= DEADLOCK_K:
                deadlock_reason = f"The same intent could not be grounded for {intent_repeat_count + 1} consecutive steps (ungroundable-intent deadlock)"
            elif optimistic_streak >= DEADLOCK_K:
                deadlock_reason = f"For {optimistic_streak} consecutive steps, assessments were optimistic but the task remained unfinished (hallucinated-completion deadlock)"
        if deadlock_reason:
            print(f"Deadlock reassessment triggered: {deadlock_reason}")
            rr = re_review_deadlock(
                state,
                filtered_admissible,
                intent_action or last_intent,
                deadlock_reason,
                client,
                model,
                max_tokens,
            )
            llm_call_count += 1
            if rr:
                selected_action = rr
                print(f"Reassessment selected: {selected_action}")
        if took_intent:
            print(
                f"Intent-direct execution: {selected_action}  (intent: {intent_action[:40]})"
            )
        elif selected_action:
            pass
        else:
            filtered_admissible = rank_by_rarity(state, filtered_admissible)
            print(
                f"Stage 1: filtering {len(filtered_admissible)} candidates using knowledge-base principles..."
            )
            candidates_to_simulate = quick_filter_candidates(
                state,
                filtered_admissible,
                target_count=10,
                client=client,
                model=model,
                max_tokens=max_tokens,
                current_think=think_text,
            )
            llm_call_count += 1
            for a in filtered_admissible[:2]:
                if a not in candidates_to_simulate:
                    candidates_to_simulate.append(a)
            print(
                f"[Diagnostic] Intent-direct execution did not match; intent_action='{intent_action}'"
            )
            print(
                f"[Diagnostic] Full executable set ({len(admissible)} actions): {admissible}"
            )
            print(
                f"[Diagnostic] Candidates selected by quick_filter ({len(candidates_to_simulate)} actions):"
            )
            for c in candidates_to_simulate:
                print(f"      · {c}")
            print(f"Stage 2: simulating {len(candidates_to_simulate)} candidates...")
            candidates = []
            for act in candidates_to_simulate:
                pred_obs = simulate_action(state, act, client, model, max_tokens)
                llm_call_count += 2
                candidates.append((act, pred_obs))
            print(f"[Diagnostic] Candidate predictions (observation + task progress):")
            for act, pred in candidates:
                lines = pred.split("\n", 1)
                obs_part = lines[0][:60] if lines else pred[:60]
                progress_part = lines[1][:60] if len(lines) > 1 else ""
                print(f"      · {act}")
                print(f"        {obs_part}{('...' if len(lines[0]) > 60 else '')}")
                if progress_part:
                    print(
                        f"        {progress_part}{('...' if len(lines[1]) > 60 else '')}"
                    )
            selected_action = select_best_action(
                state, candidates, client, model, max_tokens
            )
            llm_call_count += 1
            if not selected_action or selected_action not in [c[0] for c in candidates]:
                selected_action = candidates[0][0]
            for act, pred in candidates:
                if act == selected_action:
                    is_optimistic = (
                        "Near completion" in pred or "Task completed" in pred
                    )
                    break
        print(f"Selected action: {selected_action}")
        trajectory.append(StepRecord(kind="act", text=selected_action))
        last_actions.append(selected_action)
        admissible_before = admissible
        observation, reward, done, info = env.step([selected_action])
        observation = process_ob(observation[0])
        reward = reward[0]
        done = done[0]
        admissible_after = get_admissible(info)
        trajectory.append(StepRecord(kind="ob", text=observation))
        print(f"Observation: {observation[:80]}...")
        print(f"Reward: {reward}, done: {done}")
        state.observations.append((selected_action, observation))
        state.knowledge["tried_actions"].add(selected_action)
        update_knowledge_from_observation(
            state,
            selected_action,
            observation,
            admissible_before=admissible_before,
            admissible_after=admissible_after,
        )
        fp_before = state_fingerprint(admissible_before)
        if fp_before is not None:
            state.knowledge["visited_states"].add(fp_before)
        update_stuck_state(stuck_state, last_actions)
        if intent_action and same_intent_as_last and (not took_intent):
            intent_repeat_count += 1
        else:
            intent_repeat_count = 0
        last_intent = intent_action
        if is_optimistic and (not done):
            optimistic_streak += 1
        else:
            optimistic_streak = 0
        if deadlock_reason:
            intent_repeat_count = 0
            optimistic_streak = 0
        prompt += f"\n> think: {think_text}\n> {selected_action}\n{observation}\n"
        if done:
            print(f"Task completed! Final reward: {reward}")
            print(f"LLM calls in this episode: {llm_call_count}")
            return (reward, True, trajectory, llm_call_count)
    print("Maximum step limit reached; task incomplete")
    print(f"LLM calls in this episode: {llm_call_count}")
    return (0.0, False, trajectory, llm_call_count)


def write_results_to_excel(
    rs: List[float],
    cnts: List[int],
    out_path: str,
    method: str = "WorldModel Universal",
    llm_calls_total: int = 0,
    llm_calls_per_task: List[int] = None,
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"
    ws["A1"] = "Method"
    ws["B1"] = method
    ws["A2"] = "Total LLM Calls"
    ws["B2"] = llm_calls_total
    ws["A3"] = "Avg LLM Calls per Task"
    ws["B3"] = f"{llm_calls_total / max(1, sum(cnts)):.1f}"
    ws["A5"] = "Task Type"
    ws["B5"] = "Success"
    ws["C5"] = "Total"
    ws["D5"] = "Success Rate"
    task_types = ["put", "clean", "heat", "cool", "examine", "puttwo"]
    for i, (tt, r, c) in enumerate(zip(task_types, rs, cnts), start=6):
        ws[f"A{i}"] = tt
        ws[f"B{i}"] = int(r)
        ws[f"C{i}"] = c
        ws[f"D{i}"] = f"{r / max(1, c):.2%}"
    total_row = 6 + len(task_types)
    ws[f"A{total_row}"] = "Total"
    ws[f"B{total_row}"] = int(sum(rs))
    ws[f"C{total_row}"] = sum(cnts)
    ws[f"D{total_row}"] = f"{sum(rs) / max(1, sum(cnts)):.2%}"
    wb.save(out_path)
    print(f"Results saved to: {out_path}")


def load_prompts(prompt_file: str) -> Dict[str, str]:
    with open(prompt_file, "r", encoding="utf-8") as f:
        return json.load(f)


def main():
    config_file = "base_config.yaml"
    with open(config_file) as f:
        config = yaml.safe_load(f)
    from alfworld.agents.environment.alfred_tw_env import AlfredTWEnv

    env = AlfredTWEnv(config, train_eval="eval_out_of_distribution")
    env = env.init_env(batch_size=1)
    prompt_file = "prompts/alfworld_3prompts.json"
    prompts = load_prompts(prompt_file)
    client = build_client()
    model = os.getenv("OPENAI_MODEL", "qwen3-max-2026-01-23")
    max_tokens = int(os.getenv("OPENAI_MAX_TOKENS", "4096"))
    prefixes = {
        "pick_and_place": "put",
        "pick_clean_then_place": "clean",
        "pick_heat_then_place": "heat",
        "pick_cool_then_place": "cool",
        "look_at_obj": "examine",
        "pick_two_obj": "puttwo",
    }
    cnts = [0] * 6
    rs = [0] * 6
    llm_calls_total = 0
    llm_calls_per_task = []
    for idx in range(134):
        ob, info = env.reset()
        ob = "\n".join(ob[0].split("\n\n")[1:])
        name = "/".join(info["extra.gamefile"][0].split("/")[-3:-1])
        print(f"\n{'=' * 60}")
        print(f"Task {idx + 1}/134: {name}")
        print(f"{'=' * 60}")
        r = 0.0
        for i, (k, v) in enumerate(prefixes.items()):
            if name.startswith(k):
                prompt = (
                    "Interact with a household to solve a task. Here are two examples.\n"
                    + prompts[f"react_{v}_1"]
                    + prompts[f"react_{v}_0"]
                    + "\nHere is the task.\n"
                )
                init_prompt = prompt + ob + "\n>"
                r, done, trajectory, llm_calls = run_episode(
                    env,
                    init_prompt,
                    ob,
                    info,
                    client=client,
                    model=model,
                    max_tokens=max_tokens,
                )
                llm_calls_total += llm_calls
                llm_calls_per_task.append(llm_calls)
                rs[i] += r
                cnts[i] += 1
                break
        overall = sum(rs) / max(1, sum(cnts))
        print(f"\nCurrent progress: {idx + 1}/134")
        print(f"  r={r:.1f} rs={[int(x) for x in rs]} cnts={cnts} avg={overall:.4f}")
        print(
            f"  Total LLM calls: {llm_calls_total}, mean per task: {llm_calls_total / max(1, sum(cnts)):.1f}"
        )
    write_results_to_excel(
        rs,
        cnts,
        "test4_wm_universal_results-1step.xlsx",
        method="WorldModel Universal (1-step simulate)",
        llm_calls_total=llm_calls_total,
        llm_calls_per_task=llm_calls_per_task,
    )


if __name__ == "__main__":
    main()
